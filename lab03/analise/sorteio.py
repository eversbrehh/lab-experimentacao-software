"""Sorteio da amostra-ouro e geração das planilhas de rotulagem (seção 6 do enunciado).

`python -m analise.sorteio --fila dados/amostra/amostra_prevista.csv`

Sorteia 60 repositórios e 5 reservas da fila congelada e 5 deploys da janela de cada um,
sempre com a semente `SEMENTE`. As listas de releases vêm do cache do pipeline, lido
diretamente dos arquivos: o comando não usa o token nem a rede, e falha se a lista de
algum repositório sorteado não estiver no cache.

Saídas em `--saida` (padrão `dados/rotulagem`): `amostra_ouro.csv`, com os 65
repositórios e suas releases, e `avaliador_1.xlsx` a `avaliador_3.xlsx`, com os 60 da
amostra e as colunas de `docs/protocolo-rotulagem.md`.
"""

import argparse
import csv
import hashlib
import json
import os
import re
import sys
from datetime import date, datetime
from urllib.parse import parse_qsl, quote, urlencode, urlsplit, urlunsplit

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from metricas.janela import dentro_da_janela

SEMENTE = 42
TAMANHO_FILA = 300
N_AMOSTRA = 60
N_RESERVAS = 5
RELEASES_POR_REPOSITORIO = 5
N_AVALIADORES = 3

AMOSTRA = "amostra"
RESERVA = "reserva"
SUBSTITUIDO = "substituído"  # sorteado para a amostra, mas fora dos aprovados
DESCARTADO = "descartado"  # reserva fora dos aprovados

COLUNAS_CSV = (
    "ordem_sorteio",
    "papel",
    "full_name",
    "tag_name",
    "published_at",
    "url_release",
    "url_compare",
)
ABA_REPOSITORIOS = "repositorios"
ABA_RELEASES = "releases"
COLUNAS_REPOSITORIO = ("full_name", "url", "tipo_projeto", "entregas_reais", "observacao")
COLUNAS_RELEASE = ("full_name", "tag_name", "url_release", "url_compare", "corretiva", "observacao")
COLUNAS_DE_LINK = ("url", "url_release", "url_compare")
VALORES_ACEITOS = {
    "tipo_projeto": ("biblioteca/framework", "aplicação/serviço", "ferramenta CLI", "outro"),
    "entregas_reais": ("sim", "não", "incerto"),
    "corretiva": ("sim", "não"),
}
LARGURAS = {"full_name": 34, "tag_name": 24, "observacao": 50}
LARGURA_PADRAO = 22

API = "https://api.github.com"
SITE = "https://github.com"


class ErroDeSorteio(Exception):
    """Entrada ausente ou insuficiente para o sorteio; a mensagem diz o que corrigir."""


def main(argv: list[str] | None = None) -> int:
    args = _argumentos(argv)
    try:
        inicio, fim, pasta_cache = ler_config(args.config)
        fila = _ler_csv(args.fila)
        aprovados = set(_ler_csv(args.aprovados)["full_name"]) if args.aprovados else None
        linhas = montar_amostra(
            fila, lambda nome: ler_releases_do_cache(pasta_cache, nome), inicio, fim, aprovados
        )
        gravar_saidas(linhas, args.saida, args.sobrescrever)
    except ErroDeSorteio as erro:
        print(f"erro: {erro}", file=sys.stderr)
        return 1
    na_amostra = [linha for linha in linhas if linha["papel"] == AMOSTRA]
    print(
        f"{len({linha['full_name'] for linha in na_amostra})} repositórios e "
        f"{len(na_amostra)} releases na amostra; saídas em {args.saida}"
    )
    return 0


def _argumentos(argv: list[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="python -m analise.sorteio",
        description="Sorteia a amostra-ouro e gera as planilhas de rotulagem.",
    )
    parser.add_argument("--fila", required=True, help="CSV da fila congelada, com a coluna full_name")
    parser.add_argument("--config", default="config.yaml", help="config do pipeline (janela e pasta do cache)")
    parser.add_argument(
        "--aprovados",
        help="CSV com a coluna full_name (ex.: dados/saida/repositorios.csv); "
        "sorteado ausente dele é trocado pela primeira reserva livre",
    )
    parser.add_argument("--saida", default=os.path.join("dados", "rotulagem"), help="pasta das saídas")
    parser.add_argument(
        "--sobrescrever",
        action="store_true",
        help="grava mesmo se as saídas já existirem; apaga rótulos já preenchidos nas planilhas",
    )
    return parser.parse_args(argv)


def ler_config(caminho: str) -> tuple[date, date, str]:
    """Início e fim da janela e a pasta do cache, lidos do config do pipeline.

    A leitura é própria, e não por `pipeline.config`, porque aquele módulo exige o token.
    O caminho do cache é resolvido a partir da pasta do arquivo, como no pipeline.
    """
    try:
        with open(caminho, encoding="utf-8") as arquivo:
            config = yaml.safe_load(arquivo)
        inicio, fim = config["janela"]["inicio"], config["janela"]["fim"]
        cache = str(config["caminhos"]["cache"])
    except OSError as erro:
        raise ErroDeSorteio(f"não foi possível ler {caminho}: {erro.strerror}") from erro
    except (yaml.YAMLError, KeyError, TypeError) as erro:
        raise ErroDeSorteio(f"{caminho} não tem janela.inicio, janela.fim e caminhos.cache") from erro
    if not isinstance(inicio, date) or not isinstance(fim, date):
        raise ErroDeSorteio(f"{caminho}: as datas da janela devem estar no formato AAAA-MM-DD")
    pasta = os.path.dirname(os.path.abspath(caminho))
    return inicio, fim, os.path.normpath(os.path.join(pasta, cache))


def _ler_csv(caminho: str) -> pd.DataFrame:
    try:
        tabela = pd.read_csv(caminho)
    except (OSError, ValueError) as erro:
        raise ErroDeSorteio(f"não foi possível ler {caminho}: {erro}") from erro
    if "full_name" not in tabela.columns:
        raise ErroDeSorteio(f"{caminho} não tem a coluna full_name")
    return tabela


def sortear_repositorios(fila: pd.DataFrame, n: int = N_AMOSTRA + N_RESERVAS) -> list[str]:
    """`n` repositórios dos `TAMANHO_FILA` primeiros da fila, na ordem do sorteio.

    Os 60 primeiros de `n = 65` são os mesmos de `n = 60`: as reservas não mudam a amostra.
    """
    fila = fila.head(TAMANHO_FILA)
    if len(fila) < n:
        raise ErroDeSorteio(f"a fila tem {len(fila)} repositórios; o sorteio precisa de {n}")
    return fila.sample(n=n, random_state=SEMENTE)["full_name"].tolist()


def atribuir_papeis(sorteados: list[str], aprovados: set[str] | None = None) -> list[str]:
    """Papel de cada sorteado, na mesma ordem: os 60 primeiros são amostra, os demais, reservas.

    Com `aprovados`, quem não está no conjunto sai (`substituído` na amostra, `descartado`
    entre as reservas), e cada vaga aberta na amostra é ocupada pela primeira reserva
    aprovada ainda não usada. Não há novo sorteio.
    """
    papeis = [AMOSTRA if posicao < N_AMOSTRA else RESERVA for posicao in range(len(sorteados))]
    if aprovados is None:
        return papeis
    for posicao, nome in enumerate(sorteados):
        if nome not in aprovados:
            papeis[posicao] = SUBSTITUIDO if papeis[posicao] == AMOSTRA else DESCARTADO
    vagas = papeis.count(SUBSTITUIDO)
    for posicao, papel in enumerate(papeis):
        if vagas and papel == RESERVA:
            papeis[posicao] = AMOSTRA
            vagas -= 1
    if vagas:
        raise ErroDeSorteio(
            f"{papeis.count(SUBSTITUIDO)} sorteados saíram da amostra e as reservas aprovadas "
            f"não bastam: faltam {vagas}"
        )
    return papeis


def deploys_com_anterior(releases: list[dict], inicio: date, fim: date) -> list[tuple[dict | None, dict]]:
    """Pares `(anterior, deploy)` dos deploys da janela, em ordem de publicação.

    Mesma regra de `pipeline.releases.releases_com_commits`: deploy é release com `draft`
    e `prerelease` falsos, e a anterior é o deploy imediatamente anterior por
    `published_at`, dentro ou fora da janela (`None` na primeira release da história).
    """
    deploys = sorted(
        (r for r in releases if not r["draft"] and not r["prerelease"] and r["published_at"] is not None),
        key=lambda r: r["published_at"],
    )
    return [
        (deploys[posicao - 1] if posicao > 0 else None, deploy)
        for posicao, deploy in enumerate(deploys)
        if dentro_da_janela(deploy["published_at"], inicio, fim)
    ]


def sortear_releases(pares: list, n: int = RELEASES_POR_REPOSITORIO) -> list:
    """`n` itens de `pares`, devolvidos na ordem original.

    A semente é a mesma para todo repositório, então o sorteio de um não depende dos
    demais: trocar um repositório por uma reserva não altera as releases dos outros.
    """
    posicoes = pd.Series(range(len(pares))).sample(n=n, random_state=SEMENTE)
    return [pares[posicao] for posicao in sorted(posicoes)]


def montar_amostra(
    fila: pd.DataFrame, releases_de, inicio: date, fim: date, aprovados: set[str] | None = None
) -> list[dict]:
    """Linhas do `amostra_ouro.csv`: uma por release sorteada dos 65 repositórios.

    `releases_de(full_name)` devolve todas as releases do repositório, no formato de
    `pipeline.releases.coletar_releases`. Quem saiu (`substituído` ou `descartado`) fica
    com uma única linha, sem release: suas releases não são lidas, para que a falta delas
    no cache não impeça o sorteio dos demais.
    """
    sorteados = sortear_repositorios(fila)
    papeis = atribuir_papeis(sorteados, aprovados)
    linhas = []
    for ordem, (nome, papel) in enumerate(zip(sorteados, papeis), start=1):
        if papel in (SUBSTITUIDO, DESCARTADO):
            linhas.append({"ordem_sorteio": ordem, "papel": papel, "full_name": nome})
            continue
        pares = deploys_com_anterior(releases_de(nome), inicio, fim)
        if len(pares) < RELEASES_POR_REPOSITORIO:
            raise ErroDeSorteio(
                f"{nome} tem {len(pares)} deploys na janela; o sorteio precisa de "
                f"{RELEASES_POR_REPOSITORIO}"
            )
        for anterior, deploy in sortear_releases(pares):
            linhas.append(
                {
                    "ordem_sorteio": ordem,
                    "papel": papel,
                    "full_name": nome,
                    "tag_name": deploy["tag_name"],
                    "published_at": deploy["published_at"].isoformat(),
                    "url_release": deploy["html_url"],
                    "url_compare": _url_compare(nome, anterior, deploy),
                }
            )
    return linhas


def _url_compare(nome: str, anterior: dict | None, deploy: dict) -> str:
    if anterior is None:
        return ""
    base, head = (quote(release["tag_name"], safe="/@") for release in (anterior, deploy))
    return f"{SITE}/{nome}/compare/{base}...{head}"


def ler_releases_do_cache(pasta_cache: str, full_name: str) -> list[dict]:
    """Todas as releases do repositório, lidas das páginas gravadas por `pipeline.cliente_http`.

    Repete a regra do cache: um arquivo por URL, com o nome `sha256(URL).json` e os
    parâmetros da URL em ordem alfabética. Não faz requisições.
    """
    url = f"{API}/repos/{full_name}/releases?per_page=100"
    brutas = []
    while url:
        arquivo = os.path.join(pasta_cache, hashlib.sha256(url.encode("utf-8")).hexdigest() + ".json")
        try:
            with open(arquivo, encoding="utf-8") as aberto:
                pagina = json.load(aberto)
            corpo, link = pagina["corpo"], pagina["cabecalhos"].get("link")
        except (OSError, ValueError, KeyError, TypeError, AttributeError) as erro:
            raise ErroDeSorteio(
                f"a lista de releases de {full_name} não está completa no cache ({pasta_cache}); "
                "rode o sorteio na máquina que tem o cache da coleta"
            ) from erro
        if corpo is None:
            raise ErroDeSorteio(f"{full_name} não existe mais (404 no cache)")
        brutas.extend(corpo)
        url = _proxima_pagina(link)
    return [
        {
            "tag_name": bruta["tag_name"],
            "published_at": datetime.fromisoformat(bruta["published_at"]) if bruta.get("published_at") else None,
            "draft": bool(bruta["draft"]),
            "prerelease": bool(bruta["prerelease"]),
            "html_url": bruta["html_url"],
        }
        for bruta in brutas
    ]


def _proxima_pagina(link: str | None) -> str | None:
    achado = re.search(r'<([^>]+)>\s*;\s*rel="next"', link or "")
    if not achado:
        return None
    partes = urlsplit(achado.group(1))
    consulta = urlencode(sorted(parse_qsl(partes.query, keep_blank_values=True)))
    return urlunsplit((partes.scheme, partes.netloc, partes.path, consulta, ""))


def gravar_saidas(linhas: list[dict], pasta: str, sobrescrever: bool = False) -> None:
    """Grava `amostra_ouro.csv` e as planilhas dos avaliadores, todas com o mesmo conteúdo.

    Sem `sobrescrever`, recusa-se a gravar se alguma saída já existir: regravar uma
    planilha apagaria os rótulos preenchidos nela.
    """
    do_csv = os.path.join(pasta, "amostra_ouro.csv")
    das_planilhas = [os.path.join(pasta, f"avaliador_{n}.xlsx") for n in range(1, N_AVALIADORES + 1)]
    existentes = [caminho for caminho in (do_csv, *das_planilhas) if os.path.exists(caminho)]
    if existentes and not sobrescrever:
        raise ErroDeSorteio(
            f"{', '.join(existentes)} já existe; use --sobrescrever para gravar de novo "
            "(os rótulos já preenchidos nas planilhas são apagados)"
        )
    os.makedirs(pasta, exist_ok=True)
    with open(do_csv, "w", encoding="utf-8", newline="") as arquivo:
        escritor = csv.DictWriter(arquivo, fieldnames=COLUNAS_CSV, lineterminator="\n")
        escritor.writeheader()
        escritor.writerows(linhas)
    na_amostra = [linha for linha in linhas if linha["papel"] == AMOSTRA]
    for caminho in das_planilhas:
        gravar_planilha(caminho, na_amostra)


def gravar_planilha(caminho: str, linhas: list[dict]) -> None:
    """Planilha de um avaliador: abas `repositorios` e `releases`, com os rótulos em branco."""
    repositorios = list(dict.fromkeys(linha["full_name"] for linha in linhas))
    planilha = Workbook()
    _preencher_aba(
        planilha.active,
        ABA_REPOSITORIOS,
        COLUNAS_REPOSITORIO,
        [{"full_name": nome, "url": f"{SITE}/{nome}"} for nome in repositorios],
    )
    _preencher_aba(planilha.create_sheet(), ABA_RELEASES, COLUNAS_RELEASE, linhas)
    planilha.save(caminho)


def _preencher_aba(aba, titulo: str, colunas: tuple[str, ...], linhas: list[dict]) -> None:
    aba.title = titulo
    aba.append(colunas)
    for linha in linhas:
        aba.append([linha.get(coluna) or None for coluna in colunas])
    aba.freeze_panes = "A2"
    ultima = len(linhas) + 1
    for indice, coluna in enumerate(colunas, start=1):
        letra = aba.cell(row=1, column=indice).column_letter
        aba.cell(row=1, column=indice).font = Font(bold=True)
        aba.column_dimensions[letra].width = LARGURAS.get(coluna, LARGURA_PADRAO)
        if coluna in COLUNAS_DE_LINK:
            for celula in aba[letra][1:]:
                if celula.value:
                    celula.hyperlink = celula.value
                    celula.style = "Hyperlink"
        if coluna in VALORES_ACEITOS and ultima > 1:
            validacao = DataValidation(
                type="list",
                formula1=f'"{",".join(VALORES_ACEITOS[coluna])}"',
                allow_blank=True,
                showErrorMessage=True,  # sem isto o Excel mostra a lista, mas aceita qualquer texto
            )
            validacao.error = "Use um dos valores aceitos no protocolo de rotulagem."
            aba.add_data_validation(validacao)
            validacao.add(f"{letra}2:{letra}{ultima}")


if __name__ == "__main__":
    sys.exit(main())
