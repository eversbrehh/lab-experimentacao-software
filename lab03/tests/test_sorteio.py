import hashlib
import json
from datetime import date, datetime, timedelta, timezone

import pandas as pd
import pytest
from openpyxl import load_workbook

from analise import sorteio
from pipeline import cliente_http, releases as coleta_de_releases
from analise.sorteio import (
    AMOSTRA,
    DESCARTADO,
    RESERVA,
    SUBSTITUIDO,
    ErroDeSorteio,
    atribuir_papeis,
    deploys_com_anterior,
    ler_releases_do_cache,
    montar_amostra,
    sortear_releases,
    sortear_repositorios,
)

INICIO = date(2025, 10, 1)
FIM = date(2026, 9, 30)
BASE = datetime(2026, 1, 1, 12, tzinfo=timezone.utc)
CONFIG = """
janela:
  inicio: 2025-10-01
  fim: 2026-09-30
caminhos:
  cache: cache
  saida: saida
"""


def fila(quantidade=80):
    return pd.DataFrame({"ordem": range(1, quantidade + 1), "full_name": [f"org/repo-{i:03d}" for i in range(quantidade)]})


def lista_releases(quantidade, **campos):
    return [
        {
            "tag_name": f"v{i}",
            "published_at": BASE + timedelta(days=i),
            "draft": False,
            "prerelease": False,
            "html_url": f"https://exemplo/v{i}",
            **campos,
        }
        for i in range(quantidade)
    ]


def gravar_cache(pasta, url, corpo, link=None):
    pasta.mkdir(exist_ok=True)
    cabecalhos = {"link": link} if link else {}
    arquivo = pasta / (hashlib.sha256(url.encode("utf-8")).hexdigest() + ".json")
    arquivo.write_text(json.dumps({"url": url, "status": 200, "cabecalhos": cabecalhos, "corpo": corpo}), encoding="utf-8")


def brutas(quantidade):
    return [{**release, "published_at": release["published_at"].isoformat()} for release in lista_releases(quantidade)]


def url_releases(nome):
    return f"https://api.github.com/repos/{nome}/releases?per_page=100"


@pytest.fixture
def projeto(tmp_path):
    """Pasta com config, fila de 80 repositórios e cache com 8 deploys de cada um."""
    (tmp_path / "config.yaml").write_text(CONFIG, encoding="utf-8")
    fila().to_csv(tmp_path / "fila.csv", index=False)
    for nome in fila()["full_name"]:
        gravar_cache(tmp_path / "cache", url_releases(nome), brutas(8))
    return tmp_path


def executar(projeto, saida="rotulagem", *extras):
    return sorteio.main(
        ["--fila", str(projeto / "fila.csv"), "--config", str(projeto / "config.yaml"), "--saida", str(projeto / saida), *extras]
    )


def celulas(caminho):
    return {aba.title: [[celula.value for celula in linha] for linha in aba.iter_rows()] for aba in load_workbook(caminho)}


# sortear_repositorios


def test_semente_fica_no_codigo_e_o_sorteio_se_repete():
    assert sorteio.SEMENTE == 42
    assert sortear_repositorios(fila()) == sortear_repositorios(fila())


def test_sorteia_65_repositorios_distintos_da_fila():
    sorteados = sortear_repositorios(fila())
    assert len(sorteados) == len(set(sorteados)) == 65
    assert set(sorteados) <= set(fila()["full_name"])


def test_os_60_primeiros_de_65_sao_os_mesmos_de_60():
    assert sortear_repositorios(fila(), 65)[:60] == sortear_repositorios(fila(), 60)


def test_so_os_300_primeiros_da_fila_entram_no_sorteio():
    assert sortear_repositorios(fila(400)) == sortear_repositorios(fila(300))


def test_fila_menor_que_o_sorteio_falha():
    with pytest.raises(ErroDeSorteio, match="a fila tem 64 repositórios"):
        sortear_repositorios(fila(64))


# atribuir_papeis


def test_sem_aprovados_sao_60_da_amostra_e_5_reservas():
    assert atribuir_papeis(sortear_repositorios(fila())) == [AMOSTRA] * 60 + [RESERVA] * 5


def test_sorteado_fora_dos_aprovados_e_trocado_pela_primeira_reserva():
    sorteados = sortear_repositorios(fila())
    papeis = atribuir_papeis(sorteados, set(sorteados) - {sorteados[3]})
    assert papeis[3] == SUBSTITUIDO
    assert papeis[60:] == [AMOSTRA] + [RESERVA] * 4
    assert papeis.count(AMOSTRA) == 60


def test_reserva_que_tambem_caiu_e_pulada():
    sorteados = sortear_repositorios(fila())
    papeis = atribuir_papeis(sorteados, set(sorteados) - {sorteados[3], sorteados[60]})
    assert papeis[60:] == [DESCARTADO, AMOSTRA] + [RESERVA] * 3
    assert papeis.count(AMOSTRA) == 60


def test_reservas_insuficientes_falham():
    sorteados = sortear_repositorios(fila())
    with pytest.raises(ErroDeSorteio, match="faltam 1"):
        atribuir_papeis(sorteados, set(sorteados[6:]))


# deploys_com_anterior e sortear_releases


def test_deploys_ignoram_rascunho_e_prerelease_e_guardam_a_anterior():
    releases = lista_releases(4)
    releases[1]["prerelease"] = True
    releases[2]["draft"] = True
    pares = deploys_com_anterior(releases, INICIO, FIM)
    assert [(a and a["tag_name"], d["tag_name"]) for a, d in pares] == [(None, "v0"), ("v0", "v3")]


def test_anterior_pode_estar_fora_da_janela():
    releases = lista_releases(2)
    releases[0]["published_at"] = datetime(2025, 9, 30, tzinfo=timezone.utc)
    pares = deploys_com_anterior(releases, INICIO, FIM)
    assert [(a["tag_name"], d["tag_name"]) for a, d in pares] == [("v0", "v1")]


def test_sorteia_5_releases_em_ordem_e_se_repete():
    pares = list(range(20))
    sorteadas = sortear_releases(pares)
    assert len(set(sorteadas)) == 5
    assert sorteadas == sorted(sorteadas) == sortear_releases(pares)


# montar_amostra


def test_amostra_tem_5_releases_de_cada_um_dos_65():
    linhas = montar_amostra(fila(), lambda nome: lista_releases(8), INICIO, FIM)
    assert len(linhas) == 325
    assert [linha["ordem_sorteio"] for linha in linhas] == [ordem for ordem in range(1, 66) for _ in range(5)]
    assert sum(linha["papel"] == AMOSTRA for linha in linhas) == 300


def test_url_compare_usa_o_deploy_anterior_e_fica_vazia_na_primeira_release():
    linhas = montar_amostra(fila(), lambda nome: lista_releases(5), INICIO, FIM)
    nome = linhas[0]["full_name"]
    assert linhas[0]["tag_name"] == "v0" and linhas[0]["url_compare"] == ""
    assert linhas[1]["url_compare"] == f"https://github.com/{nome}/compare/v0...v1"
    assert linhas[1]["url_release"] == "https://exemplo/v1"


def test_tag_com_caractere_especial_e_codificada_na_url():
    releases = lista_releases(5)
    releases[0]["tag_name"] = "pkg@1.0 #1"
    linhas = montar_amostra(fila(), lambda nome: releases, INICIO, FIM)
    assert linhas[1]["url_compare"].endswith("/compare/pkg@1.0%20%231...v1")


def test_repositorio_com_menos_de_5_deploys_falha_com_mensagem_clara():
    with pytest.raises(ErroDeSorteio, match="tem 4 deploys na janela; o sorteio precisa de 5"):
        montar_amostra(fila(), lambda nome: lista_releases(4), INICIO, FIM)


def test_trocar_um_repositorio_nao_muda_as_releases_dos_demais():
    sorteados = sortear_repositorios(fila())
    antes = montar_amostra(fila(), lambda nome: lista_releases(8), INICIO, FIM)
    depois = montar_amostra(fila(), lambda nome: lista_releases(8), INICIO, FIM, set(sorteados) - {sorteados[0]})

    def dos_demais(linhas):
        return [{**linha, "papel": None} for linha in linhas if linha["full_name"] != sorteados[0]]

    assert dos_demais(antes) == dos_demais(depois)


def test_quem_saiu_fica_com_uma_linha_e_suas_releases_nao_sao_lidas():
    sorteados = sortear_repositorios(fila())
    fora = {sorteados[0], sorteados[60]}

    def releases_de(nome):
        assert nome not in fora, f"leu as releases de {nome}, que saiu da amostra"
        return lista_releases(8)

    linhas = montar_amostra(fila(), releases_de, INICIO, FIM, set(sorteados) - fora)
    assert [linha for linha in linhas if linha["full_name"] in fora] == [
        {"ordem_sorteio": 1, "papel": SUBSTITUIDO, "full_name": sorteados[0]},
        {"ordem_sorteio": 61, "papel": DESCARTADO, "full_name": sorteados[60]},
    ]
    assert sum(linha["papel"] == AMOSTRA for linha in linhas) == 300


# ler_releases_do_cache


def test_le_as_releases_do_cache_seguindo_a_paginacao(tmp_path):
    todas = brutas(3)
    proxima = "https://api.github.com/repos/org/repo/releases?per_page=100&page=2"
    gravar_cache(tmp_path, url_releases("org/repo"), todas[:2], f'<{proxima}>; rel="next"')
    gravar_cache(tmp_path, "https://api.github.com/repos/org/repo/releases?page=2&per_page=100", todas[2:])
    assert ler_releases_do_cache(str(tmp_path), "org/repo") == lista_releases(3, notas=None)  # cache sem `body`


class RespostaFalsa:
    def __init__(self, corpo, cabecalhos=None):
        self.status_code = 200
        self.headers = cabecalhos or {}
        self._corpo = corpo

    def json(self):
        return self._corpo


def test_le_o_cache_gravado_pelo_cliente_http_do_pipeline(tmp_path, monkeypatch):
    # Contrato com pipeline.cliente_http: se a regra do nome do arquivo ou o formato do
    # cache mudarem lá, este teste falha, em vez de o sorteio falhar só na execução real.
    todas = [{**bruta, "body": "notas", "assets": []} for bruta in brutas(3)]
    proxima = "https://api.github.com/repositories/1/releases?per_page=100&page=2"
    respostas = [RespostaFalsa(todas[:2], {"Link": f'<{proxima}>; rel="next"'}), RespostaFalsa(todas[2:])]
    monkeypatch.setattr(cliente_http, "_requisitar", lambda url, cabecalhos: respostas.pop(0))
    cliente_http.configurar("token-de-teste", str(tmp_path))

    coletadas = coleta_de_releases.coletar_releases("org", "repo")

    assert not respostas
    assert ler_releases_do_cache(str(tmp_path), "org/repo") == coletadas == lista_releases(3, notas="notas")


def test_rascunho_sem_data_fica_com_published_at_none(tmp_path):
    gravar_cache(tmp_path, url_releases("org/repo"), [{**brutas(1)[0], "published_at": None, "draft": True}])
    assert ler_releases_do_cache(str(tmp_path), "org/repo")[0]["published_at"] is None


def test_repositorio_fora_do_cache_falha_sem_chamar_a_api(tmp_path):
    with pytest.raises(ErroDeSorteio, match="org/repo não está completa no cache"):
        ler_releases_do_cache(str(tmp_path), "org/repo")


def test_pagina_seguinte_ausente_falha(tmp_path):
    proxima = "https://api.github.com/repos/org/repo/releases?page=2&per_page=100"
    gravar_cache(tmp_path, url_releases("org/repo"), brutas(2), f'<{proxima}>; rel="next"')
    with pytest.raises(ErroDeSorteio, match="não está completa no cache"):
        ler_releases_do_cache(str(tmp_path), "org/repo")


def test_404_no_cache_falha(tmp_path):
    gravar_cache(tmp_path, url_releases("org/repo"), None)
    with pytest.raises(ErroDeSorteio, match="não existe mais"):
        ler_releases_do_cache(str(tmp_path), "org/repo")


# main e saídas


def test_gera_o_csv_e_as_tres_planilhas(projeto):
    assert executar(projeto) == 0
    saida = projeto / "rotulagem"
    assert sorted(p.name for p in saida.iterdir()) == [
        "amostra_ouro.csv",
        "avaliador_1.xlsx",
        "avaliador_2.xlsx",
        "avaliador_3.xlsx",
    ]
    tabela = pd.read_csv(saida / "amostra_ouro.csv", keep_default_na=False)
    assert list(tabela.columns) == list(sorteio.COLUNAS_CSV)
    assert len(tabela) == 325
    assert tabela.groupby("papel")["full_name"].nunique().to_dict() == {AMOSTRA: 60, RESERVA: 5}


def test_planilha_tem_as_colunas_do_protocolo_e_nenhuma_da_heuristica(projeto):
    executar(projeto)
    abas = celulas(projeto / "rotulagem" / "avaliador_1.xlsx")
    assert list(abas) == ["repositorios", "releases"]
    assert abas["repositorios"][0] == ["full_name", "url", "tipo_projeto", "entregas_reais", "observacao"]
    assert abas["releases"][0] == ["full_name", "tag_name", "url_release", "url_compare", "corretiva", "observacao"]
    assert len(abas["repositorios"]) == 61
    assert len(abas["releases"]) == 301
    assert all(linha[2:] == [None, None, None] for linha in abas["repositorios"][1:])
    assert all(linha[4:] == [None, None] for linha in abas["releases"][1:])


def test_planilha_tem_links_diretos_e_listas_de_valores_aceitos(projeto):
    executar(projeto)
    planilha = load_workbook(projeto / "rotulagem" / "avaliador_2.xlsx")
    repositorios, releases = planilha["repositorios"], planilha["releases"]
    nome = repositorios["A2"].value
    assert repositorios["B2"].hyperlink.target == f"https://github.com/{nome}"
    assert releases["C3"].hyperlink.target == releases["C3"].value
    assert releases["D3"].hyperlink.target.startswith(f"https://github.com/{releases['A3'].value}/compare/")
    listas = {str(v.sqref): v.formula1 for aba in (repositorios, releases) for v in aba.data_validations.dataValidation}
    assert listas == {
        "C2:C61": '"biblioteca/framework,aplicação/serviço,ferramenta CLI,outro"',
        "D2:D61": '"sim,não,incerto"',
        "E2:E301": '"sim,não"',
    }


def test_listas_recusam_valor_fora_dos_aceitos(projeto):
    executar(projeto)
    planilha = load_workbook(projeto / "rotulagem" / "avaliador_3.xlsx")
    validacoes = [v for aba in planilha for v in aba.data_validations.dataValidation]
    assert len(validacoes) == 3
    assert all(v.showErrorMessage for v in validacoes)


def test_duas_execucoes_produzem_saidas_identicas(projeto):
    executar(projeto, "a")
    executar(projeto, "b")
    assert (projeto / "a" / "amostra_ouro.csv").read_bytes() == (projeto / "b" / "amostra_ouro.csv").read_bytes()
    for n in (1, 2, 3):
        assert celulas(projeto / "a" / f"avaliador_{n}.xlsx") == celulas(projeto / "b" / f"avaliador_{n}.xlsx")
    assert celulas(projeto / "a" / "avaliador_1.xlsx") == celulas(projeto / "a" / "avaliador_3.xlsx")


def test_nao_sobrescreve_planilhas_existentes_sem_o_argumento(projeto, capsys):
    executar(projeto)
    planilha = projeto / "rotulagem" / "avaliador_1.xlsx"
    antes = planilha.read_bytes()
    assert executar(projeto) == 1
    assert "já existe; use --sobrescrever" in capsys.readouterr().err
    assert planilha.read_bytes() == antes
    assert executar(projeto, "rotulagem", "--sobrescrever") == 0


def test_aprovados_trocam_o_sorteado_pela_reserva_nas_planilhas(projeto):
    sorteados = sortear_repositorios(fila())
    aprovados = projeto / "repositorios.csv"
    pd.DataFrame({"full_name": [nome for nome in fila()["full_name"] if nome != sorteados[0]]}).to_csv(aprovados, index=False)
    assert executar(projeto, "rotulagem", "--aprovados", str(aprovados)) == 0
    nomes = [linha[0] for linha in celulas(projeto / "rotulagem" / "avaliador_1.xlsx")["repositorios"][1:]]
    assert nomes == sorteados[1:61]


def test_sorteado_que_saiu_nao_precisa_estar_no_cache(projeto):
    sorteados = sortear_repositorios(fila())
    aprovados = projeto / "repositorios.csv"
    pd.DataFrame({"full_name": [nome for nome in fila()["full_name"] if nome != sorteados[0]]}).to_csv(aprovados, index=False)
    arquivo = hashlib.sha256(url_releases(sorteados[0]).encode("utf-8")).hexdigest() + ".json"
    (projeto / "cache" / arquivo).unlink()
    assert executar(projeto, "rotulagem", "--aprovados", str(aprovados)) == 0
    tabela = pd.read_csv(projeto / "rotulagem" / "amostra_ouro.csv", keep_default_na=False)
    assert len(tabela) == 321
    assert tabela[tabela["full_name"] == sorteados[0]].iloc[0].to_dict() == {
        "ordem_sorteio": 1,
        "papel": SUBSTITUIDO,
        "full_name": sorteados[0],
        "tag_name": "",
        "published_at": "",
        "url_release": "",
        "url_compare": "",
    }


def test_repositorio_fora_do_cache_encerra_com_mensagem_e_sem_saidas(projeto, capsys):
    for arquivo in (projeto / "cache").iterdir():
        arquivo.unlink()
    assert executar(projeto) == 1
    assert "não está completa no cache" in capsys.readouterr().err
    assert not (projeto / "rotulagem").exists()


def test_config_ou_fila_invalidos_encerram_com_mensagem(projeto, capsys):
    (projeto / "fila.csv").write_text("nome\nx\n", encoding="utf-8")
    assert executar(projeto) == 1
    assert "não tem a coluna full_name" in capsys.readouterr().err
    (projeto / "config.yaml").write_text("janela: {}\n", encoding="utf-8")
    assert executar(projeto) == 1
    assert "não tem janela.inicio" in capsys.readouterr().err
